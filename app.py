# -*- coding: utf-8 -*-
"""
ProveMax - Servidor local simple (Flask)
- Sirve la pagina de la tienda
- Expone el catalogo de productos (productos.json)
- Registra los pedidos en pedidos.xlsx

Uso:
    pip install flask
    python app.py
    -> abrir http://localhost:5000
"""

import json
import os
from datetime import datetime

from flask import Flask, jsonify, render_template, request

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    EXCEL_OK = True
except ImportError:
    EXCEL_OK = False

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PRODUCTOS_FILE = os.path.join(BASE_DIR, "productos.json")
PEDIDOS_FILE = os.path.join(BASE_DIR, "pedidos.xlsx")

app = Flask(__name__, template_folder="templates")


# ----------------------------------------------------------------------
# Catalogo
# ----------------------------------------------------------------------
def cargar_productos():
    with open(PRODUCTOS_FILE, "r", encoding="utf-8") as f:
        return json.load(f)


def buscar_producto(productos, pid):
    for p in productos:
        if p["id"] == pid:
            return p
    return None


# ----------------------------------------------------------------------
# Excel: pedidos
# ----------------------------------------------------------------------
def excel_disponible():
    return EXCEL_OK


def asegurar_excel():
    """Crea el archivo Excel con cabeceras si no existe."""
    if not EXCEL_OK:
        return None
    if os.path.exists(PEDIDOS_FILE):
        return load_workbook(PEDIDOS_FILE)
    wb = Workbook()
    ws = wb.active
    ws.title = "Pedidos"
    headers = [
        "N Pedido", "Fecha", "Cliente", "Producto", "Categoria",
        "Unidad", "Cantidad", "Precio Unit.", "Subtotal", "Total Pedido", "Estado",
    ]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="0B5ED7")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    anchos = [10, 18, 18, 24, 14, 16, 10, 13, 12, 14, 12]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[chr(64 + i)].width = ancho
    wb.save(PEDIDOS_FILE)
    return wb


def registrar_pedido(items, cliente="Cliente web"):
    """
    items: lista de {"id": int, "qty": int}
    Devuelve dict con numero de pedido, total y detalle.
    """
    productos = cargar_productos()
    fecha = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    numero = generar_numero_pedido()

    lineas = []
    total = 0.0
    for it in items:
        p = buscar_producto(productos, it.get("id"))
        qty = int(it.get("qty", 0))
        if not p or qty <= 0:
            continue
        subtotal = round(p["price"] * qty, 2)
        total += subtotal
        lineas.append({
            "producto": p["name"],
            "categoria": p["category"],
            "unidad": p["unit"],
            "cantidad": qty,
            "precio": p["price"],
            "subtotal": subtotal,
        })
    total = round(total, 2)

    if not lineas:
        return {"ok": False, "error": "El pedido no contiene productos validos."}

    wb = asegurar_excel()
    if wb is None:
        return {
            "ok": False,
            "error": "openpyxl no esta instalado (pip install openpyxl).",
        }

    ws = wb["Pedidos"]
    for ln in lineas:
        ws.append([
            numero, fecha, cliente, ln["producto"], ln["categoria"],
            ln["unidad"], ln["cantidad"], ln["precio"], ln["subtotal"],
            total, "Pendiente",
        ])
    wb.save(PEDIDOS_FILE)

    return {
        "ok": True,
        "numero": numero,
        "fecha": fecha,
        "total": total,
        "items": len(lineas),
        "detalle": lineas,
    }


def generar_numero_pedido():
    """Genera numero correlativo P-0001, P-0002... leyendo filas previas."""
    if not os.path.exists(PEDIDOS_FILE):
        return "P-0001"
    try:
        wb = load_workbook(PEDIDOS_FILE, read_only=True)
        ws = wb["Pedidos"]
        ultimo = 0
        for i, fila in enumerate(ws.iter_rows(min_row=2, max_col=1, values_only=True)):
            val = fila[0] if fila else None
            if val and str(val).startswith("P-"):
                try:
                    ultimo = max(ultimo, int(str(val).split("-")[1]))
                except ValueError:
                    pass
        wb.close()
        return f"P-{ultimo + 1:04d}"
    except Exception:
        return "P-0001"


# ----------------------------------------------------------------------
# Rutas
# ----------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/productos", methods=["GET"])
def api_productos():
    return jsonify(cargar_productos())


@app.route("/api/pedidos", methods=["POST"])
def api_pedidos():
    data = request.get_json(silent=True) or {}
    items = data.get("items", [])
    cliente = data.get("cliente", "Cliente web").strip() or "Cliente web"
    resultado = registrar_pedido(items, cliente)
    codigo = 200 if resultado.get("ok") else 400
    return jsonify(resultado), codigo


@app.route("/api/pedidos", methods=["GET"])
def api_pedidos_listar():
    """Lista breve de pedidos registrados (util para revisar rapidamente)."""
    if not os.path.exists(PEDIDOS_FILE) or not EXCEL_OK:
        return jsonify({"ok": True, "pedidos": []})
    wb = load_workbook(PEDIDOS_FILE, read_only=True)
    ws = wb["Pedidos"]
    vistos = {}
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not fila or not fila[0]:
            continue
        numero = fila[0]
        if numero not in vistos:
            vistos[numero] = {"numero": numero, "fecha": fila[1], "cliente": fila[2], "total": fila[9]}
    wb.close()
    return jsonify({"ok": True, "pedidos": list(vistos.values())})


if __name__ == "__main__":
    print("=" * 55)
    print(" ProveMax - servidor local")
    print(" Abre en tu navegador: http://localhost:5000")
    print(f" Pedidos se guardan en: {PEDIDOS_FILE}")
    if not EXCEL_OK:
        print(" AVISO: instala openpyxl ->  pip install openpyxl")
    print("=" * 55)
    app.run(host="127.0.0.1", port=5000, debug=True)
